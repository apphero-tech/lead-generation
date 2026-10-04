from openpyxl import load_workbook

from leadgen.db import connect, set_step
from leadgen.export import CONTACT_COLUMNS, export_state


def test_export_columns_and_rows(tmp_path):
    conn = connect(tmp_path / "t.db")
    conn.execute("INSERT INTO institutions (unitid, state, name, website, sector, system_name, is_system) VALUES ('1','FL','Test U','https://t.edu','Public','',0)")
    conn.execute(
        "INSERT INTO persons (id, unitid, person_key, first_name, last_name, email, email_status, phone_status, "
        "responsibilities, manual_check, check_reasons) VALUES (1,'1','jane|doe','Jane','Doe','jane.doe@t.edu',"
        "'deduced','not found','x',1,'email deduced')"
    )
    conn.execute("INSERT INTO person_roles VALUES (1,'adv_vp','VP Advancement','exact','https://t.edu/a',NULL)")
    conn.execute("INSERT INTO person_roles VALUES (1,'fdn_ed','VP Advancement','exact','https://t.edu/a',NULL)")
    set_step(conn, "1", "extract", "done")
    conn.commit()
    path = export_state(conn, "FL", tmp_path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Contacts", "Coverage", "Stats", "Read me"]
    ws = wb["Contacts"]
    header = [c.value for c in ws[1]]
    assert header == CONTACT_COLUMNS
    row = dict(zip(header, [c.value for c in ws[2]]))
    assert row["email_status"] == "deduced" and row["manual_check_needed"] == "yes"
    assert row["target_profiles"] == "VP Advancement; Executive Director, University Foundation"
    assert ws.max_row == 2  # one row per person, not per profile


def test_linkedin_search_link(tmp_path):
    from leadgen.export import linkedin_search_url
    url = linkedin_search_url("Amy", "Layman", "Keiser University-Ft Lauderdale")
    assert url == "https://www.linkedin.com/search/results/people/?keywords=Amy+Layman+Keiser+University"
    conn = connect(tmp_path / "t.db")
    conn.execute("INSERT INTO institutions (unitid, state, name, website, sector, system_name, is_system) "
                 "VALUES ('1','FL','Test U','https://t.edu','Public','',0)")
    conn.execute("INSERT INTO persons (id, unitid, person_key, first_name, last_name, email_status, phone_status, "
                 "manual_check) VALUES (1,'1','jane|doe','Jane','Doe','not found','not found',0)")
    conn.execute("INSERT INTO person_roles VALUES (1,'cio','CIO','exact','https://t.edu/a',NULL)")
    conn.commit()
    ws = load_workbook(export_state(conn, "FL", tmp_path))["Contacts"]
    cell = ws.cell(row=2, column=CONTACT_COLUMNS.index("linkedin_search") + 1)
    assert cell.value == "Open LinkedIn search" and "keywords=Jane+Doe+Test+U" in cell.hyperlink.target


def test_first_columns_and_bold_name(tmp_path):
    assert CONTACT_COLUMNS[:3] == ["name", "email", "linkedin_search"]
    assert CONTACT_COLUMNS.count("email") == 1 and "first_name" in CONTACT_COLUMNS
    conn = connect(tmp_path / "t.db")
    conn.execute("INSERT INTO institutions (unitid, state, name, website, sector, system_name, is_system) "
                 "VALUES ('1','FL','Test U','https://t.edu','Public','',0)")
    conn.execute("INSERT INTO persons (id, unitid, person_key, first_name, last_name, email, email_status, "
                 "phone_status, manual_check) VALUES (1,'1','jane|doe','Jane','Doe','jd@t.edu','published','not found',0)")
    conn.execute("INSERT INTO person_roles VALUES (1,'cio','CIO','exact','https://t.edu/a',NULL)")
    conn.commit()
    ws = load_workbook(export_state(conn, "FL", tmp_path))["Contacts"]
    assert [ws.cell(2, i).value for i in (1, 2)] == ["Doe, Jane", "jd@t.edu"]
    assert ws.cell(2, 1).font.bold and ws.cell(2, 1).font.size == 13
