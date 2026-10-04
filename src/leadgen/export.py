"""Step 4: per-state Excel export (Contacts, Coverage, Stats, Read me sheets)."""
from __future__ import annotations

from collections import Counter
import re
from pathlib import Path
from urllib.parse import urlencode
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .profiles import PROFILE_BY_ID, PROFILES

# The columns used most come first; details follow.
CONTACT_COLUMNS = [
    "last_name", "first_name", "email", "email_status", "linkedin_search", "phone", "institution",
    "target_profiles", "found_title", "responsibilities",
    "state", "institution_website", "system", "sector", "unit", "contact_level", "role_family",
    "email_source", "phone_status", "source_urls", "last_verified_date",
    "manual_check_needed", "check_reasons",
]
NAME_COLUMNS = ("last_name", "first_name")
NAME_FONT = Font(bold=True, size=13)
LINK_COLUMNS = {"linkedin_search"}


def linkedin_search_url(first: str, last: str, institution: str) -> str:
    """People-search link for a manual check on LinkedIn (the tool itself never visits LinkedIn).
    The campus suffix is dropped ("Keiser University-Ft Lauderdale" -> "Keiser University")."""
    core = re.split(r"\s*[-–]\s*", institution or "")[0]
    return "https://www.linkedin.com/search/results/people/?" + urlencode({"keywords": f"{first} {last} {core}".strip()})

README = [
    ("What this file is", "Contacts found on public university web pages for the target profiles. Free sources only."),
    ("email_status = published", "CERTAIN. The address was written on an official web page next to the person's name (see email_source)."),
    ("email_status = deduced", "GUESSED. Not published. Built from the institution's email format (e.g. first.last@), learned from other published addresses. Must be checked by hand."),
    ("email_status = not found", "No published address and no reliable email format for that institution."),
    ("phone_status = published", "Phone number printed next to the person on the source page (may be an office line)."),
    ("unit", "Faculty, school or office the person belongs to (from the source page title)."),
    ("email published via directory", "email_source starting with 'official directory search' = found by searching the person's name in the institution's own online directory."),
    ("found_title", "The exact title on the source page. target_profiles is the profile it was matched to."),
    ("manual_check_needed = yes", "Something needs a human look; check_reasons says what (deduced email, closest-match title, old source, several candidates, ...)."),
    ("last_verified_date", "Date the tool last read the source page online."),
    ("responsibilities", "'(from source page)' = sentence taken from the page; '(typical scope for this title)' = generic description of the job."),
    ("linkedin_search", "Click to open a LinkedIn people search for this person (name + institution) in your browser, to check their current role by hand. The tool itself never visits LinkedIn."),
    ("contact_level", "'this campus' = found on this institution's own pages (or IPEDS names this campus's own chief). 'network-wide' = the institution is one campus of a chain sharing one website (e.g. 20 'Arizona College of Nursing' campuses); the person was found on the shared site, so is probably at headquarters or another campus."),
    ("Coverage sheet", "One row per institution: who was found for each of the 18 profiles. Empty = nobody found (never invented)."),
    ("phone_status = main switchboard (IPEDS)", "The institution's general phone number from the federal IPEDS directory, not a direct line."),
    ("closest title", "check_reasons 'this is the closest title at this institution' = nobody holds the target job; this is the nearest role (common at small schools)."),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _in(unitids: Optional[List[str]], col: str = "i.unitid") -> Tuple[str, list]:
    if unitids is None:
        return "", []
    return f" AND {col} IN ({','.join('?' * len(unitids)) or 'NULL'})", list(unitids)


def collect_rows(conn, state: str, unitids: Optional[List[str]] = None) -> List[dict]:
    extra, args = _in(unitids)
    persons = conn.execute(
        "SELECT p.*, i.name AS inst_name, i.website, i.system_name, i.sector FROM persons p "
        "JOIN institutions i ON i.unitid = p.unitid WHERE i.state = ?" + extra +
        " ORDER BY i.is_system DESC, i.name, p.last_name, p.first_name",
        [state] + args,
    ).fetchall()
    order = {p.id: n for n, p in enumerate(PROFILES)}
    rows = []
    for p in persons:
        roles = conn.execute(
            "SELECT profile_id, found_title, match_quality, source_url FROM person_roles WHERE person_id = ?",
            (p["id"],),
        ).fetchall()
        pids = sorted({r["profile_id"] for r in roles}, key=lambda x: order[x])
        titles = list(dict.fromkeys(r["found_title"] for r in roles if r["found_title"]))
        sources = list(dict.fromkeys(r["source_url"] for r in roles))
        if p["email_source"] and p["email_source"].startswith("http") and p["email_source"] not in sources:
            sources.append(p["email_source"])
        rows.append({
            "state": state,
            "institution": p["inst_name"],
            "institution_website": p["website"],
            "system": p["system_name"],
            "sector": p["sector"],
            "unit": p["unit"],
            "contact_level": p["level"] or "this campus",
            "first_name": p["first_name"],
            "last_name": p["last_name"],
            "target_profiles": "; ".join(PROFILE_BY_ID[x].target_title for x in pids),
            "found_title": " / ".join(titles),
            "role_family": "; ".join(dict.fromkeys(PROFILE_BY_ID[x].family for x in pids)),
            "responsibilities": p["responsibilities"],
            "email": p["email"],
            "email_status": p["email_status"],
            "email_source": p["email_source"],
            "phone": p["phone"],
            "phone_status": p["phone_status"],
            "source_urls": "\n".join(sources),
            "last_verified_date": p["last_verified_date"],
            "manual_check_needed": "yes" if p["manual_check"] else "no",
            "check_reasons": p["check_reasons"],
            "linkedin_search": linkedin_search_url(p["first_name"], p["last_name"], p["inst_name"]),
            "_pids": pids,
            "_unitid": p["unitid"],
        })
    return rows


def compute_stats(conn, state: str, rows: List[dict], unitids: Optional[List[str]] = None) -> List[tuple]:
    extra, args = _in(unitids)
    insts = conn.execute("SELECT unitid, is_system FROM institutions i WHERE state = ?" + extra,
                         [state] + args).fetchall()
    steps = dict(
        (r["unitid"], r["status"]) for r in conn.execute(
            "SELECT s.unitid, s.status FROM step_status s JOIN institutions i ON i.unitid = s.unitid "
            "WHERE i.state = ? AND s.step = 'extract'" + extra, [state] + args)
    )
    ex = conn.execute(
        "SELECT COALESCE(SUM(raw_candidates),0) AS raw, COALESCE(SUM(persons),0) AS kept, "
        "COALESCE(SUM(pages_scanned),0) AS pages FROM extraction_stats e "
        "JOIN institutions i ON i.unitid = e.unitid WHERE i.state = ?" + extra, [state] + args
    ).fetchone()
    n = len(rows) or 1
    email = Counter(r["email_status"] for r in rows)
    phones = sum(1 for r in rows if r["phone_status"] == "published")
    manual = sum(1 for r in rows if r["manual_check_needed"] == "yes")
    multi = sum(1 for r in rows if len(r["_pids"]) > 1)
    out = [
        ("Institutions in selection (incl. system offices)", len(insts)),
        ("Institutions processed", sum(1 for s in steps.values() if s == "done")),
        ("Institutions skipped (no website)", sum(1 for s in steps.values() if s == "skipped")),
        ("Institutions failed", sum(1 for s in steps.values() if s == "failed")),
        ("Institutions not processed yet", len(insts) - len(steps)),
        ("Pages scanned", ex["pages"]),
        ("Contacts (unique people)", len(rows)),
        ("Raw title matches before de-duplication", ex["raw"]),
        ("People covering 2+ profiles (merged into one row)", multi),
        ("Email published", f"{email['published']} ({email['published'] / n:.0%})"),
        ("Email deduced", f"{email['deduced']} ({email['deduced'] / n:.0%})"),
        ("Email not found", f"{email['not found']} ({email['not found'] / n:.0%})"),
        ("Phone published", f"{phones} ({phones / n:.0%})"),
        ("Manual check needed", f"{manual} ({manual / n:.0%})"),
        ("Cost (paid APIs)", "$0 - free sources only"),
        ("", ""),
        ("Contacts per target profile", ""),
    ]
    per = Counter(pid for r in rows for pid in r["_pids"])
    for p in PROFILES:
        out.append(("  " + p.target_title, per.get(p.id, 0)))
    return out


def _sheet(ws, header: List[str], data: List[list], widths: Dict[str, int]) -> None:
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for row in data:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 18)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for i, h in enumerate(header, 1):
        if h in LINK_COLUMNS:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                if cell.value:
                    cell.hyperlink = cell.value
                    cell.value = "Open LinkedIn search"
                    cell.font = Font(color="0563C1", underline="single")


def export_state(conn, state: str, out_dir: Path, unitids: Optional[List[str]] = None,
                 filename: Optional[str] = None) -> Path:
    """Write the Excel file for a state, or only for the given institutions."""
    rows = collect_rows(conn, state, unitids)
    extra, args = _in(unitids)
    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"
    widths = {"last_name": 18, "first_name": 16, "email": 32, "email_status": 14, "phone": 18, "institution": 30, "responsibilities": 50, "source_urls": 50, "check_reasons": 50,
              "found_title": 40, "target_profiles": 35, "email": 30, "email_source": 40,
              "institution_website": 28, "sector": 22, "system": 25, "unit": 35, "linkedin_search": 22, "contact_level": 24}
    _sheet(ws, CONTACT_COLUMNS, [[r[c] for c in CONTACT_COLUMNS] for r in rows], widths)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=len(NAME_COLUMNS)):
        for cell in row:
            cell.font = NAME_FONT
    ws.freeze_panes = "C2"  # header row and name columns stay visible while scrolling

    cov = wb.create_sheet("Coverage")
    insts = conn.execute(
        "SELECT i.unitid, i.name, i.website, i.sector, s.status, s.detail FROM institutions i "
        "LEFT JOIN step_status s ON s.unitid = i.unitid AND s.step = 'extract' "
        "WHERE i.state = ?" + extra + " ORDER BY i.is_system DESC, i.name", [state] + args
    ).fetchall()
    by_inst: Dict[str, Dict[str, List[str]]] = {}
    for r in rows:
        for pid in r["_pids"]:
            by_inst.setdefault(r["_unitid"], {}).setdefault(pid, []).append(f"{r['first_name']} {r['last_name']}")
    header = ["institution", "website", "sector", "status", "profiles_found"] + [p.target_title for p in PROFILES]
    data = []
    for i in insts:
        found = by_inst.get(i["unitid"], {})
        status = i["status"] or "not processed yet"
        if status == "skipped":
            status = "skipped - no website (manual check needed)"
        crawl = conn.execute("SELECT detail FROM step_status WHERE unitid = ? AND step = 'crawl'",
                             (i["unitid"],)).fetchone()
        if crawl and (crawl["detail"] or "").startswith("0 pages:"):
            status += " - " + crawl["detail"][9:]
        data.append([i["name"], i["website"], i["sector"], status, f"{len(found)}/{len(PROFILES)}"]
                    + [", ".join(found.get(p.id, [])) for p in PROFILES])
    _sheet(cov, header, data, {"institution": 34, "website": 28, "status": 22})

    st = wb.create_sheet("Stats")
    _sheet(st, ["metric", "value"], [list(x) for x in compute_stats(conn, state, rows, unitids)], {"metric": 50, "value": 22})

    rd = wb.create_sheet("Read me")
    _sheet(rd, ["item", "meaning"], [list(x) for x in README], {"item": 28, "meaning": 110})

    target = out_dir / state / (filename or f"contacts_{state}.xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return target
